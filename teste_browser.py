#!/usr/bin/env python3
import requests
from io import BytesIO
import openpyxl
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

url = "https://governosp.sharepoint.com/:x:/r/teams/SECGOVERNO-SECOM_Data/Shared%20Documents/001_SicomData/Miscelaneous/testeConectorPython.xlsx?d=w02cbd2a2c5a24fc9835ffd9fa1c8bbaa&csf=1&web=1&e=GVgTTE"

# Simular navegador real
headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
}

print("[*] Testando com headers de navegador real...")
response = requests.get(url, headers=headers, verify=False, timeout=30, allow_redirects=True)

print(f"[*] Status: {response.status_code}")
print(f"[*] Content-Type: {response.headers.get('content-type', 'unknown')}")
print(f"[*] Tamanho: {len(response.content)} bytes")
print(f"[*] URL final: {response.url}")

if response.status_code == 200:
    try:
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        headers = [str(cell.value) if cell.value else "" for cell in ws[1]]
        
        print(f"\n[OK] Arquivo Excel válido!")
        print(f"[OK] Cabeçalhos encontrados: {headers}")
        
        row_count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and any(cell is not None for cell in row):
                row_count += 1
        
        print(f"[OK] Total de linhas: {row_count}")
        print("\n✓ SUCESSO! Link é acessível via browser!")
        
    except Exception as e:
        print(f"[ERRO] {type(e).__name__}: {str(e)}")
else:
    print(f"[ERRO] Falha - Status {response.status_code}")
