#!/usr/bin/env python3
import requests
from io import BytesIO
import openpyxl
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

url = "https://governosp.sharepoint.com/:x:/r/teams/SECGOVERNO-SECOM_Data/Shared%20Documents/001_SicomData/Miscelaneous/testeConectorPython.xlsx?d=w02cbd2a2c5a24fc9835ffd9fa1c8bbaa&csf=1&web=1&e=GVgTTE"

# Criar sessão com cookies persistentes
session = requests.Session()

# Headers de navegador real
headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, application/octet-stream',
    'Accept-Language': 'pt-BR,pt;q=0.9,en-US;q=0.8,en;q=0.7',
    'Accept-Encoding': 'gzip, deflate, br',
    'Connection': 'keep-alive',
    'Upgrade-Insecure-Requests': '1',
    'Sec-Fetch-Dest': 'document',
    'Sec-Fetch-Mode': 'navigate',
    'Sec-Fetch-Site': 'none',
    'Cache-Control': 'max-age=0',
}

session.headers.update(headers)

print("[*] Testando com sessão de navegador...")

try:
    # Primeiro, fazer uma requisição GET para tentar obter cookies da sessão
    print("[*] Fazendo primeira requisição para estabelecer sessão...")
    response1 = session.get("https://governosp.sharepoint.com/", verify=False, timeout=10, allow_redirects=True)
    print(f"[*] Primeira resposta: Status {response1.status_code}")
    print(f"[*] Cookies da sessão: {session.cookies.get_dict()}")
    
    # Agora tentar baixar o arquivo
    print("\n[*] Tentando baixar arquivo com sessão estabelecida...")
    response = session.get(url, verify=False, timeout=30, allow_redirects=True)
    
    print(f"[*] Status: {response.status_code}")
    print(f"[*] Content-Type: {response.headers.get('content-type', 'unknown')}")
    print(f"[*] Tamanho: {len(response.content)} bytes")
    print(f"[*] URL final: {response.url[:80]}...")
    
    if response.status_code == 200 and len(response.content) > 1000:
        # Verificar se é um arquivo Excel válido
        try:
            wb = openpyxl.load_workbook(BytesIO(response.content))
            ws = wb.active
            headers_list = [str(cell.value) if cell.value else "" for cell in ws[1]]
            
            print(f"\n[OK] Arquivo Excel válido!")
            print(f"[OK] Cabeçalhos: {headers_list}")
            
            row_count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and any(cell is not None for cell in row):
                    row_count += 1
            
            print(f"[OK] Total de linhas: {row_count}")
            print("\n✓✓✓ SUCESSO! Sessão com cookies funcionou!")
            
        except Exception as e:
            print(f"[ERRO] Não é Excel válido: {type(e).__name__}")
            print(f"[DEBUG] Primeiros 200 bytes: {response.content[:200]}")
    else:
        print(f"[ERRO] Resposta inválida (status={response.status_code}, tamanho={len(response.content)})")
        if len(response.text) < 500:
            print(f"[DEBUG] Conteúdo: {response.text}")
        else:
            print(f"[DEBUG] Conteúdo (truncado): {response.text[:500]}")
            
except Exception as e:
    print(f"[ERRO] Exceção: {type(e).__name__}")
    print(f"[ERRO] Mensagem: {str(e)}")
